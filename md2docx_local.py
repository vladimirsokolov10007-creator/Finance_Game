# -*- coding: utf-8 -*-
"""Markdown -> docx converter for the two Finni project documents."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
CODE_RE = re.compile(r'`([^`]+)`')

def add_runs(par, text):
    # split on **bold** and `code`
    pos = 0
    for m in re.finditer(r'(\*\*.+?\*\*|`.+?`)', text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith('**'):
            r = par.add_run(tok[2:-2]); r.bold = True
        else:
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])

def is_table_sep(line):
    return bool(re.match(r'^\s*\|?[\s:|-]+\|?\s*$', line)) and '-' in line

def convert(md_path, docx_path):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(11)

    lines = open(md_path, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith('```'):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1  # skip closing fence
            for bl in buf:
                p = doc.add_paragraph()
                r = p.add_run(bl if bl else ' ')
                r.font.name = 'Consolas'; r.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(0)
            continue

        # table
        if stripped.startswith('|') and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            header = [c.strip() for c in stripped.strip('|').split('|')]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(header))
            t.style = 'Table Grid'
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.paragraphs[0].text = ''
                r = cell.paragraphs[0].add_run(re.sub(r'\*\*|`', '', h)); r.bold = True
            for ri, row in enumerate(rows):
                for j, c in enumerate(row):
                    if j < len(header):
                        cell = t.rows[ri + 1].cells[j]
                        cell.paragraphs[0].text = ''
                        add_runs(cell.paragraphs[0], c)
            continue

        m = re.match(r'^(#{1,6})\s+(.*)', stripped)
        if m:
            level = min(len(m.group(1)), 4)
            h = doc.add_heading('', level=level)
            add_runs(h, m.group(2))
            i += 1
            continue

        if re.match(r'^[-*]\s+', stripped):
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, re.sub(r'^[-*]\s+', '', stripped))
            i += 1
            continue

        if re.match(r'^\d+\.\s+', stripped):
            p = doc.add_paragraph(style='List Number')
            add_runs(p, re.sub(r'^\d+\.\s+', '', stripped))
            i += 1
            continue

        if stripped == '---':
            i += 1
            continue

        if stripped == '':
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    # footer page numbers
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    fld1 = run._r.makeelement('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar', {'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType': 'begin'})
    instr = run._r.makeelement(qn('w:instrText'), {qn('xml:space'): 'preserve'})
    instr.text = 'PAGE'
    fld2 = run._r.makeelement('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar', {'{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType': 'end'})
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)

    doc.save(docx_path)
    print('saved', docx_path)

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
